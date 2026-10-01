"""Caratteri di controllo invisibili: gli export Excel li sostituiscono invece di fallire."""
import io
import zipfile

from openpyxl import Workbook, load_workbook

import models
from services import excel_safety
from services.excel_export import build_full_backup_zip_bytes, build_language_workbook

excel_safety.install()

# riferimento incollato da Word: a capo manuale (\x0b) più un controllo spurio (\x02)
DIRTY = "Pancheva (2021: 3).\x0bMorphosyntactic variation\x02"
CLEAN = "Pancheva (2021: 3). Morphosyntactic variation "


def _seed_dirty_example(db_session):
    lang = models.Language(id="BUL", name_full="Bulgarian", position=1)
    db_session.add(lang)
    db_session.add(models.ParameterDef(id="NQP", position=1, name="Numeral quantification", is_active=True))
    db_session.add(models.Question(
        id="NQP_Qa", parameter_id="NQP", text="Is it marked?", is_stop_question=False, is_active=True,
    ))
    db_session.flush()
    answer = models.Answer(language_id="BUL", question_id="NQP_Qa", response_text="yes", status="approved")
    db_session.add(answer)
    db_session.flush()
    db_session.add(models.Example(answer_id=answer.id, number="1", textarea="dva stola", reference=DIRTY))
    db_session.commit()
    return lang


def _all_values(wb):
    return {v for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row}


def test_clean_text_replaces_forbidden_characters():
    assert excel_safety.clean_text(DIRTY) == CLEAN


def test_clean_text_keeps_tab_and_line_breaks():
    assert excel_safety.clean_text("a\tb\nc\r\nd") == "a\tb\nc\r\nd"


def test_workbook_cleans_dirty_text_even_after_double_install():
    # se una nuova versione di openpyxl cambia il metodo patchato, fallisce qui
    excel_safety.install()
    excel_safety.install()
    ws = Workbook().active
    ws.append([DIRTY])
    assert ws["A1"].value == CLEAN


def test_language_workbook_with_dirty_reference(db_session):
    lang = _seed_dirty_example(db_session)
    wb = build_language_workbook(db_session, lang, is_admin=True)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    assert CLEAN in _all_values(load_workbook(buf))


def test_full_backup_with_dirty_reference(db_session):
    lang = _seed_dirty_example(db_session)
    data = build_full_backup_zip_bytes(db_session, [lang])
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        wb = load_workbook(io.BytesIO(zf.read("languages/BUL.xlsx")))
    assert CLEAN in _all_values(wb)
